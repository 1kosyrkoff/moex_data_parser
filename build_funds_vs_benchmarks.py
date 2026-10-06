#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Отчёт «Фонды ВИМ против своих бенчмарков» — пересборка на новой выгрузке.

    python build_funds_vs_benchmarks.py \
        --template funds_vs_benchmarks_02102026.xlsx \
        --data all_funds_perform_DDMMYYYY.xlsx \
        [--out funds_vs_benchmarks_DDMMYYYY.xlsx] [--cny cbr_cny.csv] [--recalc]

Шаблон — прошлая версия отчёта. Из неё берутся оформление ячеек, графики, привязка
фонд → бенчмарк, примечания и курс CNY/RUB за уже покрытые даты. Новые даты курса ЦБ
подтягиваются с cbr.ru (XML_dynamic, R01375) или из --cny: CSV «effective_date;rate»
(дата вступления курса в силу, как в файле ЦБ; запятая в числах допускается).

Методика (как в исходной книге):
  * уровень = 1 + накопленная доходность выгрузки; RU000A0JT4S1 с 14.08.2026 × 100 (дробление 1:100);
  * прирост за период = уровень на конец / уровень на начало − 1; граница — последнее значение
    на дату или раньше (допуск 10 дней); «2 года» = максимальное окно, если история ≥ 90% периода;
  * дневные приросты — только между соседними строками, где оба значения есть;
  * TE = STDEV(дневная избыточная)·√252, бета = cov/var по тем же парам, просадки — по уровням;
  * композитный ранг периода = средний из рангов доли обогнавших, средней и медианной альфы;
    три лучших периода получают отдельные листы.
Формулы E:K листов фондов, сводки и листов периодов — живые; бета, просадки и число
наблюдений пишутся значениями. Без --recalc формулы досчитает Excel при открытии.
"""
from __future__ import annotations

import argparse
import json
import math
import os
import posixpath
import re
import shutil
import subprocess
import sys
import tempfile
import urllib.request
import warnings
import zipfile
from copy import copy
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from lxml import etree
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font
from openpyxl.styles.cell_style import StyleArray
from pandas.tseries.offsets import BDay, DateOffset

warnings.filterwarnings("ignore")

# ============================================================================ настройки
SPLITS = [("RU000A0JT4S1", "2026-08-14", 100.0)]       # (ISIN, дата, множитель) — дробления паёв
CBR_CNY_ID = "R01375"
PERIODS = [("1 месяц", 1), ("2 месяца", 2), ("3 месяца", 3), ("6 месяцев", 6), ("12 месяцев", 12),
           ("2 года", 24), ("3 года", 36), ("4 года", 48), ("5 лет", 60)]
STALE_DAYS = 10
MAXWIN_COVERAGE = 0.90
YUAN_SHEET, YUAN_ISIN, YUAN_RAW = "Ликвидность. Юань", "RU000A107D41", "RUSFARCNY"
INNOV_SHEET = "Инновационный"
RANTIE = [("Облигации. Рантье", "RU000A1079N7", "RUCBTRNS"),
          ("Сбалансированный. Рантье", "RU000A10AV80", "VTBFG0017")]
BLUE_FONT, RED_FONT = "FF1F5FA9", "FFB3372F"
HI_FILL = "FFFFF3C4"
BLUE, RED, LIGHT = "2A78D6", "E34948", "9BB9DD"
MINUS, EMPTY = "\u2212", '""'
MATRIX_FIRST_ROW, PERIOD_TABLE_ROW, FUND_TABLE_ROW, PERIOD_SHEET_FIRST_ROW = 46, 31, 7, 6


# ============================================================================ расчёт
@dataclass
class PeriodDef:
    label: str
    months: int
    nominal_start: pd.Timestamp
    start: pd.Timestamp | None
    end: pd.Timestamp
    status: str          # ok | maxwin | short


@dataclass
class Metrics:
    status: str
    start: pd.Timestamp | None = None
    end: pd.Timestamp | None = None
    start_pos: int | None = None
    end_pos: int | None = None
    f_pos: int | None = None
    b_pos: int | None = None
    f_end: int | None = None
    b_end: int | None = None
    days: int | None = None
    E: float = float("nan")
    F: float = float("nan")
    G: float = float("nan")
    H: float = float("nan")
    I: float = float("nan")
    J: float = float("nan")
    K: float | str = float("nan")
    L: float = float("nan")
    M: float = float("nan")
    N: float = float("nan")
    O: int = 0


def period_defs(dates: pd.DatetimeIndex) -> list[PeriodDef]:
    end, first = dates.max(), dates.min()
    out = []
    for label, m in PERIODS:
        nom = end - DateOffset(months=m)
        if nom >= first:
            out.append(PeriodDef(label, m, nom, dates[dates <= nom].max(), end, "ok"))
        elif (end - first).days / (end - nom).days >= MAXWIN_COVERAGE:
            out.append(PeriodDef(label, m, nom, first, end, "maxwin"))
        else:
            out.append(PeriodDef(label, m, nom, None, end, "short"))
    return out


def asof_pos(s: pd.Series, pos: int, dates: pd.DatetimeIndex) -> int | None:
    d0 = dates[pos]
    for p in range(pos, -1, -1):
        if (d0 - dates[p]).days > STALE_DAYS:
            return None
        if not pd.isna(s.iat[p]):
            return p
    return None


def daily_ret(s: pd.Series) -> pd.Series:
    return s / s.shift(1) - 1


def mdd(x: pd.Series) -> float:
    x = x.dropna()
    return float((x / x.cummax() - 1).min()) if len(x) else float("nan")


def fund_metrics(lv: pd.DataFrame, fcode: str, bcode: str | None, p: PeriodDef) -> Metrics:
    if bcode is None:
        return Metrics("nobench")
    if p.status == "short":
        return Metrics("short")
    dates, f, b = lv.index, lv[fcode], lv[bcode]
    end_pos = dates.get_loc(p.end)
    start_pos = int(np.argmax((f.notna() & b.notna()).values)) if p.status == "maxwin" else dates.get_loc(p.start)
    fp, bp = asof_pos(f, start_pos, dates), asof_pos(b, start_pos, dates)
    fe, be = asof_pos(f, end_pos, dates), asof_pos(b, end_pos, dates)
    if None in (fp, bp, fe, be) or fe <= fp or be <= bp:
        return Metrics("nodata")
    m = Metrics(p.status, dates[start_pos], dates[end_pos], start_pos, end_pos, fp, bp, fe, be)
    m.days = (m.end - m.start).days
    m.E = f.iat[fe] / f.iat[fp] - 1
    m.F = b.iat[be] / b.iat[bp] - 1
    m.G = m.E - m.F
    m.H = (1 + m.E) / (1 + m.F) - 1
    m.I = (1 + m.E) ** (365.25 / m.days) - (1 + m.F) ** (365.25 / m.days)
    w = daily_ret(f).iloc[start_pos + 1:end_pos + 1]
    x = daily_ret(b).iloc[start_pos + 1:end_pos + 1]
    ok = w.notna() & x.notna()
    m.O = int(ok.sum())
    m.J = float((w - x)[ok].std(ddof=1) * math.sqrt(252)) if m.O > 1 else float("nan")
    m.K = m.I / m.J if m.J and not math.isnan(m.J) else ""
    if m.O > 1:
        m.L = float(np.cov(w[ok].values, x[ok].values, ddof=1)[0, 1] / np.var(x[ok].values, ddof=1))
    m.M = mdd(f.iloc[start_pos:end_pos + 1])
    m.N = mdd(b.iloc[start_pos:end_pos + 1])
    return m


def period_scores(pdefs, res, compared):
    rows = []
    for i, p in enumerate(pdefs):
        al = [res[s][i].G for s in compared if res[s][i].status in ("ok", "maxwin")]
        if p.status == "short" or not al:
            rows.append(None)
            continue
        k = sum(a > 0 for a in al)
        rows.append(dict(label=p.label, n=len(al), k=k, share=k / len(al), mean=float(np.mean(al)),
                         median=float(np.median(al))))
    valid = [r for r in rows if r]
    ranks = [pd.Series([r[key] for r in valid]).rank(ascending=False, method="average")
             for key in ("share", "mean", "median")]
    for j, r in enumerate(valid):
        r["composite"] = float(sum(rk.iloc[j] for rk in ranks) / 3)
    return rows


# ============================================================================ данные
def fetch_cbr(d1: pd.Timestamp, d2: pd.Timestamp) -> pd.Series:
    url = ("https://www.cbr.ru/scripts/XML_dynamic.asp?date_req1=%s&date_req2=%s&VAL_NM_RQ=%s"
           % (d1.strftime("%d/%m/%Y"), d2.strftime("%d/%m/%Y"), CBR_CNY_ID))
    with urllib.request.urlopen(url, timeout=30) as r:
        return parse_cbr_xml(r.read())


def parse_cbr_xml(xml: bytes) -> pd.Series:
    root = etree.fromstring(xml)
    data = {}
    for rec in root.findall("Record"):
        val = float(rec.findtext("Value").replace(",", ".")) / float(rec.findtext("Nominal").replace(",", "."))
        data[pd.Timestamp(pd.to_datetime(rec.get("Date"), format="%d.%m.%Y"))] = val
    return pd.Series(data, dtype=float).sort_index()


def read_cny_csv(path: str) -> pd.Series:
    df = pd.read_csv(path, sep=";", decimal=",", dtype=str)
    dates = pd.to_datetime(df.iloc[:, 0].str.strip(), dayfirst=True)
    vals = df.iloc[:, 1].str.strip().str.replace(",", ".", regex=False).astype(float)
    return pd.Series(vals.values, index=pd.DatetimeIndex(dates)).sort_index()


def build_levels(raw_path: str, tmpl_data: pd.DataFrame, cny_csv: str | None):
    raw = pd.read_excel(raw_path)
    raw = raw.rename(columns={raw.columns[0]: "Дата"}).set_index("Дата").sort_index()
    raw.index = pd.DatetimeIndex(raw.index)
    lv = 1.0 + raw
    info = {"splits": []}
    for isin, d, k in SPLITS:
        d = pd.Timestamp(d)
        if isin in lv and lv.index.min() < d <= lv.index.max():
            pre = lv.loc[:d - pd.Timedelta(days=1), isin].dropna().iloc[-1]
            post = lv.loc[d:, isin].dropna().iloc[0]
            lv.loc[d:, isin] *= k
            info["splits"].append(dict(isin=isin, date=d, k=k, pre=float(pre), post_raw=float(post)))
    jumps = lv.ffill().pct_change().abs()
    big = [(c, d.date()) for c in lv.columns for d in jumps.index[jumps[c] > 0.5]]
    if big:
        print("ВНИМАНИЕ: дневные изменения > 50% (новое дробление?):", big)

    # CNY/RUB: прошлый отчёт → cbr.ru / CSV; курс сдвинут на 1 рабочий день назад
    old = tmpl_data["CNYRUB"].dropna() if "CNYRUB" in tmpl_data else pd.Series(dtype=float)
    need = [d for d in lv.index if d not in old.index]
    shifted = pd.Series(dtype=float)
    if need:
        if cny_csv:
            eff = read_cny_csv(cny_csv)
            src = f"файл {Path(cny_csv).name}"
        else:
            try:
                eff = fetch_cbr(min(need) - pd.Timedelta(days=10), max(need) + pd.Timedelta(days=10))
            except Exception as e:                     # нет сети / cbr.ru недоступен
                raise SystemExit(f"Не удалось скачать курс CNY/RUB с cbr.ru ({e}). "
                                 "Скачайте динамику курса юаня с cbr.ru и передайте файлом --cny "
                                 "(CSV: effective_date;rate).")
            src = "cbr.ru"
        shifted = pd.Series(eff.values, index=eff.index - BDay(1)).sort_index()
        info["cny_src"] = src
    cny = []
    for d in lv.index:
        if d in old.index:
            cny.append(float(old.loc[d]))
        else:
            v = shifted.asof(d)
            if pd.isna(v) or (d - shifted.index[shifted.index <= d].max()).days > STALE_DAYS:
                raise SystemExit(f"Нет курса CNY/RUB ЦБ на {d.date()} — передайте --cny")
            cny.append(float(v))
    info["cny_new"] = len(need)
    info["cny_new_from"] = min(need) if need else None
    lv["CNYRUB"] = cny
    lv["RUSFARCNY_RUB"] = lv["RUSFARCNY"] * lv["CNYRUB"] / lv["CNYRUB"].iloc[0]
    return lv, info


# ============================================================================ шаблон
def read_template_meta(wb):
    dws = wb["Данные"]
    hdr = {dws.cell(1, c).value: openpyxl.utils.get_column_letter(c)
           for c in range(1, dws.max_column + 1) if dws.cell(1, c).value}
    col2code = {v: k for k, v in hdr.items()}
    funds = []
    for ws in wb.worksheets:
        a3 = ws["A3"].value
        if not (isinstance(a3, str) and a3.startswith("ISIN ") and ws["T1"].value == "Фонд (100)"):
            continue
        m = re.match(r"ISIN (\S+) · (\S+) · Бенчмарк: (\S+) — (.*)", a3)
        fcol = re.search(r"Данные!\$([A-Z]+)\$2", ws["T2"].value).group(1)
        u2 = ws["U2"].value
        bcol = re.search(r"Данные!\$([A-Z]+)\$2", u2).group(1) if isinstance(u2, str) and u2.startswith("=") else None
        funds.append(dict(sheet=ws.title, name=ws["A1"].value, isin=m.group(1), ftype=m.group(2), bench=m.group(3),
                          fcode=col2code[fcol], bcode=col2code.get(bcol) if bcol else None))
    period_sheets = sorted([n for n in wb.sheetnames if re.match(r"^\d+\. ", n)], key=lambda n: int(n.split(".")[0]))
    ps0 = wb[period_sheets[0]]
    n_rows = 0
    while ps0.cell(PERIOD_SHEET_FIRST_ROW + n_rows, 2).value:
        n_rows += 1
    notes = {}
    for n in period_sheets:
        w = wb[n]
        for r in range(PERIOD_SHEET_FIRST_ROW, PERIOD_SHEET_FIRST_ROW + n_rows):
            if w.cell(r, 14).value:
                notes[w.cell(r, 2).value] = w.cell(r, 14).value
    sws = wb["Сводка"]
    matrix = []
    r = MATRIX_FIRST_ROW
    while sws.cell(r, 1).value:
        matrix.append(sws.cell(r, 1).value)
        r += 1
    return funds, period_sheets, n_rows, notes, matrix, hdr


# ============================================================================ форматирование
def dt(x):
    return pd.Timestamp(x).strftime("%d.%m.%Y")


def ru(x, nd):
    return f"{abs(x):.{nd}f}".replace(".", ",")


def ru_pct(x, nd=1, plus=False):
    return (MINUS if x < 0 else ("+" if plus and x > 0 else "")) + ru(x * 100, nd) + "%"


def ru_num(x, nd=2):
    return (MINUS if x < 0 else "") + ru(x, nd)


def dot_pct(x):
    return f"{x * 100:.2f}%"


def qname(sheet):
    return sheet if re.fullmatch(r"\w+", sheet) else f"'{sheet}'"


def set_font_color(cell, rgb):
    f = copy(cell.font)
    cell.font = Font(name=f.name, sz=f.sz, b=f.b, i=f.i, u=f.u, strike=f.strike, vertAlign=f.vertAlign,
                     color=rgb, charset=f.charset, family=f.family, scheme=f.scheme)


def fill_rgb(cell):
    return cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type else None


# ============================================================================ книга
def build_workbook(args, lv, info, tmpl_data):
    wb = openpyxl.load_workbook(args.template)
    funds, tmpl_period_sheets, n_tmpl_rows, notes, matrix, hdr = read_template_meta(wb)
    by_sheet = {f["sheet"]: f for f in funds}
    by_name = {f["name"]: f for f in funds}
    compared = [by_name[n]["sheet"] for n in matrix]
    assert all(by_sheet[s]["bcode"] for s in compared)

    pdefs = period_defs(lv.index)
    res = {f["sheet"]: [fund_metrics(lv, f["fcode"], f["bcode"], p) for p in pdefs] for f in funds}
    scores = period_scores(pdefs, res, compared)
    valid = [i for i, s in enumerate(scores) if s]
    top3 = sorted(valid, key=lambda i: (scores[i]["composite"], -scores[i]["mean"], i))[:len(tmpl_period_sheets)]
    top_names = {pi: f"{k}. {pdefs[pi].label}" for k, pi in enumerate(top3, 1)}
    first, end = lv.index.min(), lv.index.max()
    end_row = len(lv) + 1
    blank = StyleArray()                       # стиль «по умолчанию» для очищаемых ячеек
    data_name = Path(args.data).name

    # --- тексты
    i3m = [p.label for p in pdefs].index("3 месяца")
    i12 = [p.label for p in pdefs].index("12 месяцев")
    i2y = [p.label for p in pdefs].index("2 года")
    fy, by_, rawy = lv[YUAN_ISIN], lv["RUSFARCNY_RUB"], lv[YUAN_RAW]
    wf, wbr, wr = daily_ret(fy), daily_ret(by_), daily_ret(rawy)
    vol_f, vol_b = (s.dropna().std() * math.sqrt(252) for s in (wf, wbr))
    ok = wf.notna() & wbr.notna()
    d_corr = float(np.corrcoef(wf[ok], wbr[ok])[0, 1])
    d_beta = float(np.cov(wf[ok], wbr[ok])[0, 1] / np.var(wbr[ok], ddof=1))
    both = fy.notna() & by_.notna()
    try:
        mf, mb = (s[both].resample("ME").last().pct_change() for s in (fy, by_))
    except ValueError:                         # pandas < 2.2: месячный конец периода — "M"
        mf, mb = (s[both].resample("M").last().pct_change() for s in (fy, by_))
    okm = mf.notna() & mb.notna()
    m_corr = float(np.corrcoef(mf[okm], mb[okm])[0, 1])
    m_beta = float(np.cov(mf[okm], mb[okm])[0, 1] / np.var(mb[okm], ddof=1))
    a3 = (fund_metrics(lv, YUAN_ISIN, YUAN_RAW, pdefs[i3m]).G, res[YUAN_SHEET][i3m].G)
    a2 = (fund_metrics(lv, YUAN_ISIN, YUAN_RAW, pdefs[i2y]).G, res[YUAN_SHEET][i2y].G)
    yuan_start = res[YUAN_SHEET][i2y].start
    yuan_note = (
        "Бенчмарк переведён в рубли: RUSFARCNY (юаневая ставка) умножен на относительное изменение курса CNY/RUB. "
        "Курс ЦБ РФ сдвинут на 1 рабочий день назад с даты вступления в силу. После перевода годовая волатильность "
        f"бенчмарка {ru_pct(vol_b)} против {ru_pct(vol_f)} у пая, корреляция месячных приростов {ru_num(m_corr)}, "
        f"бета {ru_num(m_beta)} — сравнение стало сопоставимым. ВНИМАНИЕ: дневные TE, IR и бета занижены/завышены "
        "из-за рассинхрона фиксинга ЦБ и времени оценки пая "
        f"(дневная корреляция {ru_num(d_corr)}, дневная бета {ru_num(d_beta)}); смотреть на альфу за период, "
        "а не на дневные риск-метрики.")
    notes[YUAN_SHEET] = yuan_note
    days_total, yrs_total = (end - first).days, (end - first).days / 365.25
    texts = {
        "A2": f"Конец всех периодов — {dt(end)} (последняя дата в данных). Периоды: 1, 2, 3, 6, 12 месяцев и 2, 3, 4, 5 лет до этой даты.",
        "A3": (f"Источники: {data_name} (накопленные доходности паёв и индексов) · Copy of Фонды и индексы.xlsx, лист Sheet2 "
               "(привязка фонд → бенчмарк) · cny.xlsx и cbr.ru (официальный курс CNY/RUB ЦБ РФ)."),
        "A6": (f"• Ряд каждого фонда и индекса приведён к уровню: уровень = 1 + накопленная доходность из выгрузки "
               f"(база 1,000000 на {dt(first)}). Лист «Данные»."),
        "A17": (f"• ГЛУБИНА ИСТОРИИ. Выгрузка покрывает {dt(first)} – {dt(end)} — {days_total} дней ({ru(yrs_total, 2)} года). "
                "Периоды 3, 4 и 5 лет посчитать нельзя: они помечены «нет данных»."),
        "A18": (f"• Период «2 года» заменён на максимальное доступное окно {dt(first)} – {dt(end)} ({ru(yrs_total, 2)} г вместо 2,00 г). "
                "Это не полные два года — сопоставлять с внешними «2г» нельзя."
                + (f" У «{YUAN_SHEET}» окно начинается {dt(yuan_start)}: ряд {YUAN_RAW} в выгрузке начинается с этой даты."
                   if not pd.isna(yuan_start) and yuan_start != first else "")),
        "A21": ("• ВАЛЮТА — ИСПРАВЛЕНО. «Ликвидность. Юань» (RU000A107D41): пай рублёвый, а раскрытый бенчмарк RUSFARCNY — "
                "ставка в юанях. Бенчмарк переведён в рубли: RUSFARCNY_RUB = RUSFARCNY x (курс CNY/RUB на дату / курс на "
                f"{dt(first)}). Курс — официальный ЦБ РФ, сдвинут на 1 рабочий день назад с даты вступления в силу (файл "
                f"датирован днём вступления, вт–сб). Контроль: волатильность бенчмарка выросла с ~0% до {ru_pct(vol_b)} против "
                f"{ru_pct(vol_f)} у пая, корреляция месячных приростов {ru_num(m_corr)}, месячная бета {ru_num(m_beta)}. "
                f"Эффект перевода на альфу: за 3 месяца {ru_pct(a3[0], 2, True)} → {ru_pct(a3[1], 2, True)}, "
                f"за 2 года {ru_pct(a2[0], 2, True)} → {ru_pct(a2[1], 2, True)} (без перевода → с переводом)."),
        "A22": ("• ОГОВОРКА ПО «ЛИКВИДНОСТЬ. ЮАНЬ». Дневные TE, IR и бета по этому фонду искажены рассинхроном между "
                f"фиксингом курса ЦБ и временем оценки пая: дневная корреляция {ru_num(d_corr)} и дневная бета "
                f"{ru_num(d_beta)} против месячных {ru_num(m_corr)} и {ru_num(m_beta)}. Опираться на альфу за период, "
                "а не на дневные риск-метрики."),
        "A23": ("• СМЕНА МАНДАТА. «Инновационный» (RU000A0JS9P7) ранее — «ВТБ – Фонд Электроэнергетики». Бенчмарк MOEXIT "
                f"применим не ко всей истории; альфа {ru_pct(res[INNOV_SHEET][i12].G, 1, True)} за 12 мес и "
                f"{ru_pct(res[INNOV_SHEET][i2y].G, 1, True)} за 2 года включает эффект смены стратегии."),
    }
    for sp in info["splits"]:
        bench = lv["VTBFG0010"]
        d0 = bench.loc[:sp["date"] - pd.Timedelta(days=1)].dropna().index[-1]
        d1 = bench.loc[sp["date"]:].dropna().index[0]
        texts["A19"] = (
            f"• ДРОБЛЕНИЕ. {sp['isin']} («Накопительный резерв») {dt(sp['date'])}: уровень {ru(sp['pre'], 6)} → "
            f"{ru(sp['post_raw'], 6)}. Значения с этой даты умножены на {sp['k']:.0f} — дробление паёв 1:{sp['k']:.0f}. "
            f"После корректировки дневной прирост {ru_pct(sp['post_raw'] * sp['k'] / sp['pre'] - 1, 2)}, что "
            f"согласуется с {ru_pct(bench.loc[d1] / bench.loc[d0] - 1, 1)} у бенчмарка VTBFG0010 за те же дни.")
    if all(c in lv for _, c, _ in RANTIE) and all(b in lv for *_, b in RANTIE):
        ra = [fund_metrics(lv, fc, bc, pdefs[i2y]).G for _, fc, bc in RANTIE]
        texts["A20"] = (
            "• ВЫПЛАТЫ ПАЙЩИКАМ. «Облигации. Рантье» и «Сбалансированный. Рантье» — фонды с выплатами: в рядах видны "
            "регулярные просадки пая 2–5% в середине месяца. Бенчмарки RUCBTRNS и 50% IMOEX + 50% RUCBTRNS — индексы "
            f"ПОЛНОЙ доходности. Альфа этих двух фондов занижена на сумму выплат; {ru_pct(ra[0])} и {ru_pct(ra[1])} "
            "за 2 года не являются результатом управления. Поэтому в сравнение они не включены (отдельных листов по ним нет).")
    cny_txt = "cny.xlsx и cbr.ru"
    data_note = (
        f"Уровень пая/индекса = 1 + накопленная доходность из {data_name} (база 1,000000 на {dt(first)}; "
        f"ряд RUSFARCNY начинается {dt(lv['RUSFARCNY'].first_valid_index())}). RU000A0JT4S1 с 14.08.2026 умножен на 100 — "
        f"корректировка дробления паёв 1:100 (источник: указание заказчика). CNYRUB — официальный курс ЦБ РФ ({cny_txt}), "
        "сдвинут на 1 рабочий день назад с даты вступления в силу. "
        f"RUSFARCNY_RUB = RUSFARCNY x (CNYRUB / CNYRUB на {dt(first)}) — юаневая ставка, приведённая к рублю.")

    # --- Данные
    ws = wb["Данные"]
    cols = [ws.cell(1, c).value for c in range(2, len(lv.columns) + 2)]
    if cols != list(lv.columns):
        raise SystemExit(f"Состав рядов выгрузки отличается от шаблона:\n шаблон: {cols}\n выгрузка: {list(lv.columns)}")
    st_date, st_num = copy(ws["A3"]._style), copy(ws["B3"]._style)
    old_max = ws.max_row
    for i, (d, row) in enumerate(lv.iterrows()):
        r = i + 2
        c = ws.cell(r, 1)
        c.value, c._style = d.to_pydatetime(), copy(st_date)
        for j, v in enumerate(row.values, 2):
            c = ws.cell(r, j)
            c.value = None if pd.isna(v) else float(v)
            c._style = copy(st_num)
    for r in range(end_row + 1, old_max + 1):
        for c in range(1, len(lv.columns) + 2):
            ws.cell(r, c).value = None
            ws.cell(r, c)._style = copy(blank)
    ws["AR1"] = data_note
    COL = {code: openpyxl.utils.get_column_letter(c) for c, code in
           enumerate([ws.cell(1, c).value for c in range(1, len(lv.columns) + 2)], 1) if code}

    # --- листы фондов
    for fd in funds:
        ws = wb[fd["sheet"]]
        fc, bc = COL[fd["fcode"]], (COL[fd["bcode"]] if fd["bcode"] else None)
        old_last = ws.max_row
        last_styles = {col: copy(ws[f"{col}{old_last}"]._style) for col in "STUVWXY"}
        for r in range(2, end_row + 1):
            pt = EMPTY if r == 2 else f"T{r - 1}"
            ws[f"S{r}"] = f"=Данные!$A${r}"
            ws[f"T{r}"] = f"=IF(ISNUMBER(Данные!${fc}${r}),Данные!${fc}${r}*100,{pt})"
            if bc:
                pu = EMPTY if r == 2 else f"U{r - 1}"
                pv = EMPTY if r == 2 else f"V{r - 1}"
                ws[f"U{r}"] = f"=IF(ISNUMBER(Данные!${bc}${r}),Данные!${bc}${r}*100,{pu})"
                ws[f"V{r}"] = f"=IF(AND(ISNUMBER(Данные!${fc}${r}),ISNUMBER(Данные!${bc}${r})),T{r}/U{r}*100,{pv})"
            if r >= 3:
                ws[f"W{r}"] = f'=IF(AND(ISNUMBER(Данные!${fc}${r}),ISNUMBER(Данные!${fc}${r - 1})),T{r}/T{r - 1}-1,"")'
                if bc:
                    ws[f"X{r}"] = f'=IF(AND(ISNUMBER(Данные!${bc}${r}),ISNUMBER(Данные!${bc}${r - 1})),U{r}/U{r - 1}-1,"")'
                    ws[f"Y{r}"] = f'=IF(AND(ISNUMBER(W{r}),ISNUMBER(X{r})),W{r}-X{r},"")'
            if r > old_last:
                for col in "STUVWXY":
                    ws[f"{col}{r}"]._style = copy(last_styles[col])
        for r in range(end_row + 1, old_last + 1):
            for col in "STUVWXY":
                ws[f"{col}{r}"].value = None
                ws[f"{col}{r}"]._style = copy(blank)
        if fd["sheet"] == YUAN_SHEET:
            ws["A4"] = "⚠ " + yuan_note
        for i, p in enumerate(pdefs):
            r, m = FUND_TABLE_ROW + i, res[fd["sheet"]][i]
            is_value_row = isinstance(ws[f"G{r}"].value, str) and ws[f"G{r}"].value.startswith("=")
            if m.status not in ("ok", "maxwin"):
                if is_value_row:
                    raise SystemExit(f"{fd['sheet']}: период «{p.label}» больше не считается — нужна доработка")
                continue
            if not is_value_row:
                raise SystemExit(f"{fd['sheet']}: появился период «{p.label}» — нужна доработка оформления")
            ws[f"B{r}"], ws[f"C{r}"] = m.start.to_pydatetime(), m.end.to_pydatetime()
            ws[f"D{r}"] = f"=C{r}-B{r}"
            ws[f"E{r}"] = f"=Данные!${fc}${m.f_end + 2}/Данные!${fc}${m.f_pos + 2}-1"
            ws[f"F{r}"] = f"=Данные!${bc}${m.b_end + 2}/Данные!${bc}${m.b_pos + 2}-1"
            ws[f"G{r}"] = f"=E{r}-F{r}"
            ws[f"H{r}"] = f"=(1+E{r})/(1+F{r})-1"
            ws[f"I{r}"] = f"=(1+E{r})^(365.25/D{r})-(1+F{r})^(365.25/D{r})"
            ws[f"J{r}"] = f"=STDEV(Y{m.start_pos + 3}:Y{end_row})*SQRT(252)"
            ws[f"K{r}"] = f'=IFERROR(I{r}/J{r},"")'
            ws[f"L{r}"], ws[f"M{r}"], ws[f"N{r}"], ws[f"O{r}"] = m.L, m.M, m.N, m.O
            nts = []
            if m.status == "maxwin":
                yrs = f"{m.days / 365.25:.2f}"
                nts.append(f"история начинается {dt(first)}; фактическое окно {yrs} г вместо 2,00 г" if m.start == first
                           else f"история бенчмарка начинается {dt(m.start)}; фактическое окно {yrs} г вместо 2,00 г")
            if m.O < 65:
                nts.append("окно <65 набл.: TE/IR неустойчивы")
            if m.J < 0.005:
                nts.append("TE<0,5%: IR численно неустойчив")
            ws[f"P{r}"] = "; ".join(nts) if nts else None
            set_font_color(ws[f"G{r}"], BLUE_FONT if m.G > 0 else RED_FONT)

    # --- листы лучших периодов
    for k, tn in enumerate(tmpl_period_sheets):
        wb[tn].title = f"__tmp{k}"
    n_cmp = len(compared)
    for k, pi in enumerate(top3):
        ws = wb[f"__tmp{k}"]
        ws.title = top_names[pi]
        row_style = {par: {c: copy(ws.cell(PERIOD_SHEET_FIRST_ROW + par, c)._style) for c in range(1, 15)} for par in (0, 1)}
        p, s = pdefs[pi], scores[pi]
        ws["A1"] = f"Период {p.label}: фонды против своих бенчмарков"
        ws["A2"] = f"{dt(p.start)} – {dt(p.end)} · {(p.end - p.start).days} дн."
        ws["A3"] = (f"Обогнали бенчмарк: {s['k']} из {s['n']} ({s['share']:.0%}) · средняя альфа {dot_pct(s['mean'])} · "
                    f"медиана {dot_pct(s['median'])}")
        order = sorted(compared, key=lambda sh: -res[sh][pi].G)
        rr = FUND_TABLE_ROW + pi
        for j, sh in enumerate(order):
            r, fd = PERIOD_SHEET_FIRST_ROW + j, by_sheet[sh]
            for c in range(1, 15):
                ws.cell(r, c)._style = copy(row_style[j % 2][c])
            ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"], ws[f"E{r}"] = j + 1, fd["name"], fd["ftype"], fd["isin"], fd["bench"]
            for col, src in zip("FGHIJKL", "EFGHJKL"):
                ws[f"{col}{r}"] = f"={qname(sh)}!${src}${rr}"
            ws[f"M{r}"] = f'=IF(H{r}>0,"да","нет")'
            ws[f"N{r}"] = notes.get(fd["name"])
            set_font_color(ws[f"H{r}"], BLUE_FONT if res[sh][pi].G > 0 else RED_FONT)
        for r in range(PERIOD_SHEET_FIRST_ROW + n_cmp, PERIOD_SHEET_FIRST_ROW + max(n_tmpl_rows, n_cmp) + 1):
            for c in range(1, 15):
                ws.cell(r, c).value = None
                ws.cell(r, c)._style = copy(blank)

    # --- Сводка
    ws = wb["Сводка"]
    for k, v in texts.items():
        ws[k] = v
    rows = range(PERIOD_TABLE_ROW, PERIOD_TABLE_ROW + 6)
    hi_row = next(r for r in rows if fill_rgb(ws.cell(r, 1)) == HI_FILL)
    lo_row = next(r for r in rows if fill_rgb(ws.cell(r, 1)) != HI_FILL)
    st_hi = {c: copy(ws.cell(hi_row, c)._style) for c in range(1, 12)}
    st_lo = {c: copy(ws.cell(lo_row, c)._style) for c in range(1, 12)}
    for i in range(6):
        r, p, s = PERIOD_TABLE_ROW + i, pdefs[i], scores[i]
        for c in range(1, 12):
            ws.cell(r, c)._style = copy(st_hi[c] if i in top3 else st_lo[c])
        ws[f"B{r}"], ws[f"C{r}"] = p.start.to_pydatetime(), p.end.to_pydatetime()
        ws[f"J{r}"] = round(s["composite"], 2)
        ws[f"K{r}"] = f"да — лист «{top_names[i]}»" if i in top3 else "—"
    for i in range(6, 9):
        ws[f"K{PERIOD_TABLE_ROW + i}"] = f"нужна история с {dt(pdefs[i].nominal_start)}, данные с {dt(first)}"
    best = max(s["share"] for s in scores if s)
    best_l = [s["label"] for s in scores if s and abs(s["share"] - best) < 1e-12]
    bk = next(s for s in scores if s and s["label"] == best_l[0])
    if best < 0.5:
        ws["A41"] = (f"Ни в одном периоде доля обогнавших бенчмарк не достигает 50% — максимум {best:.0%} "
                     f"({bk['k']} из {bk['n']}) за " + " и за ".join(best_l) +
                     ". Три листа ниже — ОТНОСИТЕЛЬНО лучшие периоды, а не периоды уверенного опережения.")
    else:
        over = [s["label"] for s in scores if s and s["share"] >= 0.5]
        ws["A41"] = ("Доля обогнавших бенчмарк достигает 50% и выше за " + ", ".join(over) +
                     f" (максимум {best:.0%}, {bk['k']} из {bk['n']}). Три листа ниже — лучшие периоды по композитному рангу.")
    for k, pi in enumerate(top3):
        s = scores[pi]
        ws[f"A{94 + k}"] = top_names[pi]
        ws[f"B{94 + k}"] = (f"Все фонды за период {pdefs[pi].label}, отсортированы по альфе. Обогнали: {s['k']} из "
                            f"{s['n']} ({s['share']:.0%}), средняя альфа {dot_pct(s['mean'])}.")

    # --- заглушки графиков: тот же порядок, что в шаблоне (XML подставит transplant_charts)
    chart_count = {"Сводка": 1, **{f["sheet"]: (3 if f["bcode"] else 1) for f in funds}, **{top_names[pi]: 1 for pi in top3}}
    for w in wb.worksheets:
        w._charts, w._images = [], []
    for sheet, n in chart_count.items():
        for j in range(n):
            ch = BarChart()
            ch.add_data(Reference(wb[sheet], min_col=1, min_row=1, max_row=2))
            wb[sheet].add_chart(ch, f"A{1 + j}")
    wb.calculation.fullCalcOnLoad = True
    wb.active = 0
    wb.save(args.out)

    meta = dict(
        base=dt(first), end_row=end_row, n_cmp=n_cmp, n_tmpl_rows=n_tmpl_rows, top3=top3, chart_count=chart_count,
        top_names={top_names[pi]: tmpl_period_sheets[k] for k, pi in enumerate(top3)},
        period_labels={top_names[pi]: pdefs[pi].label for pi in top3},
        fund_alpha_pos={f["sheet"]: [bool(res[f["sheet"]][i].G > 0) for i in range(6)] for f in funds if f["bcode"]},
        period_alpha_pos={top_names[pi]: [bool(res[sh][pi].G > 0) for sh in sorted(compared, key=lambda sh: -res[sh][pi].G)]
                          for pi in top3},
        chart_first_row={f["sheet"]: int(np.argmax((lv[f["fcode"]].notna() &
                                                     (lv[f["bcode"]].notna() if f["bcode"] else True)).values)) + 2
                         for f in funds},
    )
    return meta, dict(funds=funds, compared=compared, pdefs=pdefs, res=res, scores=scores, top3=top3,
                      top_names=top_names)


# ============================================================================ графики
NS = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
      "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
      "pr": "http://schemas.openxmlformats.org/package/2006/relationships"}
C_, A_ = "{%s}" % NS["c"], "{%s}" % NS["a"]
R_ID = "{%s}id" % NS["r"]


def _rels(z, part):
    d, fn = posixpath.split(part)
    rp = posixpath.join(d, "_rels", fn + ".rels")
    if rp not in z.namelist():
        return {}
    out = {}
    for rel in etree.fromstring(z.read(rp)).findall("pr:Relationship", NS):
        t = rel.get("Target")
        out[rel.get("Id")] = t.lstrip("/") if t.startswith("/") else posixpath.normpath(posixpath.join(d, t))
    return out


def _sheets(z):
    rm = _rels(z, "xl/workbook.xml")
    return {s.get("name"): rm[s.get(R_ID)] for s in etree.fromstring(z.read("xl/workbook.xml")).find("m:sheets", NS)}


def _drawing(z, sheet_part):
    return next((p for p in _rels(z, sheet_part).values() if "/drawings/" in p), None)


def _set_dpt_color(dpt, color):
    sp = dpt.find("c:spPr", NS)
    if sp is None:
        sp = etree.Element(C_ + "spPr")
        after = None
        for t in ("idx", "invertIfNegative", "marker", "bubble3D", "explosion"):
            if dpt.find("c:" + t, NS) is not None:
                after = dpt.find("c:" + t, NS)
        (after.addnext(sp) if after is not None else dpt.insert(0, sp))
        fill = etree.SubElement(sp, A_ + "solidFill")
        etree.SubElement(fill, A_ + "srgbClr", val=color)
        ln = etree.SubElement(sp, A_ + "ln", w="0")
        etree.SubElement(ln, A_ + "noFill")
        return
    clr = sp.find("a:solidFill/a:srgbClr", NS)
    if clr is not None:
        clr.set("val", color)
    else:
        for el in sp.findall("a:solidFill", NS):
            sp.remove(el)
        fill = etree.Element(A_ + "solidFill")
        etree.SubElement(fill, A_ + "srgbClr", val=color)
        sp.insert(0, fill)


def _color_points(root, colors, keep=None):
    """Цвет каждого столбца; лишние точки (idx >= keep) удаляются, недостающие создаются."""
    ser = root.find(".//c:ser", NS)
    dpts = ser.findall("c:dPt", NS)
    if dpts:
        anchor = dpts[0].getprevious()
    else:
        anchor = None
        for t in ("order", "tx", "spPr", "invertIfNegative", "pictureOptions"):
            el = ser.find("c:" + t, NS)
            if el is not None:
                anchor = el
    have = {}
    for d in dpts:
        ser.remove(d)
        i = int(d.find("c:idx", NS).get("val"))
        if keep is None or i < keep:
            have[i] = d
    for i in range(len(colors)):
        if i not in have:
            d = etree.Element(C_ + "dPt")
            etree.SubElement(d, C_ + "idx", val=str(i))
            etree.SubElement(d, C_ + "invertIfNegative", val="0")
            etree.SubElement(d, C_ + "bubble3D", val="0")
            have[i] = d
    for i in sorted(have, reverse=True):
        anchor.addnext(have[i])
    for i, col in enumerate(colors):
        _set_dpt_color(have[i], col)
    if keep is not None:
        dl = ser.find("c:dLbls", NS)
        if dl is not None:
            for lbl in dl.findall("c:dLbl", NS):
                if int(lbl.find("c:idx", NS).get("val")) >= keep:
                    dl.remove(lbl)


def transplant_charts(template, src, dst, meta):
    zt, zs = zipfile.ZipFile(template), zipfile.ZipFile(src)
    ts, ss = _sheets(zt), _sheets(zs)
    parts = {}
    for sheet, n in meta["chart_count"].items():
        tsheet = meta["top_names"].get(sheet, sheet)
        is_period = sheet in meta["top_names"]
        tdraw, sdraw = _drawing(zt, ts[tsheet]), _drawing(zs, ss[sheet])
        trels, srels = _rels(zt, tdraw), _rels(zs, sdraw)
        droot = etree.fromstring(zt.read(tdraw))
        anchors = [e for e in droot if e.tag.endswith("Anchor")]
        if not (len(anchors) == n == len(srels)):
            raise SystemExit(f"{sheet}: графиков в шаблоне {len(anchors)}, ожидалось {n}")
        for k, an in enumerate(anchors, 1):
            ch = an.find(".//c:chart", NS)
            croot = etree.fromstring(zt.read(trels[ch.get(R_ID)]))
            ch.set(R_ID, f"rId{k}")
            for tag in ("numCache", "strCache"):
                for el in list(croot.iter(C_ + tag)):
                    el.getparent().remove(el)
            fs, ts_ = list(croot.iter(C_ + "f")), list(croot.iter(A_ + "t"))
            if sheet == "Сводка":
                _color_points(croot, [BLUE if i in meta["top3"] else LIGHT for i in range(6)])
            elif is_period:
                last = PERIOD_SHEET_FIRST_ROW - 1 + meta["n_cmp"]
                for f in fs:
                    t = f.text.replace(f"'{tsheet}'!", f"'{sheet}'!")
                    f.text = re.sub(r"\$([A-Z]+)\$%d:\$([A-Z]+)\$\d+" % PERIOD_SHEET_FIRST_ROW,
                                    lambda mm: f"${mm.group(1)}${PERIOD_SHEET_FIRST_ROW}:${mm.group(2)}${last}", t)
                for t in ts_:
                    if t.text and t.text.startswith("Альфа за "):
                        t.text = f"Альфа за {meta['period_labels'][sheet]}, по убыванию"
                _color_points(croot, [BLUE if p else RED for p in meta["period_alpha_pos"][sheet]], keep=meta["n_cmp"])
                shift = meta["n_cmp"] - meta["n_tmpl_rows"]
                for tag in ("from", "to"):
                    row = an.find(f"xdr:{tag}/xdr:row", NS)
                    row.text = str(int(row.text) + shift)
            else:
                fr = meta["chart_first_row"][sheet]
                for f in fs:
                    f.text = re.sub(r"\$([S-V])\$\d+:\$([S-V])\$\d+",
                                    lambda mm: f"${mm.group(1)}${fr}:${mm.group(2)}${meta['end_row']}", f.text)
                for t in ts_:
                    if t.text:
                        t.text = re.sub(r"\d{2}\.\d{2}\.\d{4} = 100", f"{meta['base']} = 100", t.text)
                if croot.find(".//c:barChart", NS) is not None:
                    _color_points(croot, [BLUE if p else RED for p in meta["fund_alpha_pos"][sheet]])
            parts[srels[f"rId{k}"]] = etree.tostring(croot, xml_declaration=True, encoding="UTF-8", standalone=True)
        parts[sdraw] = etree.tostring(droot, xml_declaration=True, encoding="UTF-8", standalone=True)
    tmp = dst + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
        for item in zs.infolist():
            data = parts.pop(item.filename, None)
            zo.writestr(item, data if data is not None else zs.read(item.filename))
    zs.close()
    os.replace(tmp, dst)
    assert not parts, parts


# ============================================================================ пересчёт (LibreOffice)
MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub RecalculateAndSave()
      ThisComponent.calculateAll()
      ThisComponent.store()
      ThisComponent.close(True)
    End Sub
</script:module>"""


def find_soffice():
    for c in (shutil.which("soffice"), "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if c and Path(c).exists():
            return c
    return None


def recalc(path: str) -> bool:
    soffice = find_soffice()
    if not soffice:
        print("LibreOffice не найден — формулы досчитает Excel при открытии.")
        return False
    env = dict(os.environ, SAL_USE_VCLPLUGIN="svp")
    with tempfile.TemporaryDirectory() as prof:
        url = Path(prof).as_uri()
        subprocess.run([soffice, "--headless", "--terminate_after_init", f"-env:UserInstallation={url}"],
                       capture_output=True, timeout=180, env=env)
        mdir = Path(prof) / "user" / "basic" / "Standard"
        if not mdir.exists():
            print("Не удалось подготовить профиль LibreOffice — пересчёт пропущен.")
            return False
        (mdir / "Module1.xba").write_text(MACRO, encoding="utf-8")
        before = os.stat(path).st_mtime_ns
        subprocess.run([soffice, "--headless", "--norestore", f"-env:UserInstallation={url}",
                        "vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application",
                        str(Path(path).absolute())], capture_output=True, timeout=600, env=env)
        ok = os.stat(path).st_mtime_ns != before
        print("Пересчёт LibreOffice:", "выполнен" if ok else "НЕ выполнен")
        return ok


def check(path, ctx):
    """Сверка пересчитанных значений с расчётом pandas."""
    wbv = openpyxl.load_workbook(path, data_only=True)
    bad, errs = [], 0
    for f in ctx["funds"]:
        if not f["bcode"]:
            continue
        ws = wbv[f["sheet"]]
        for i in range(6):
            m = ctx["res"][f["sheet"]][i]
            for col in "EFGHIJ":
                v = ws[f"{col}{FUND_TABLE_ROW + i}"].value
                if not isinstance(v, (int, float)) or abs(v - getattr(m, col)) > 1e-9:
                    bad.append((f["sheet"], col, v, getattr(m, col)))
    for w in wbv.worksheets:
        for row in w.iter_rows():
            for c in row:
                if isinstance(c.value, str) and re.match(r"^#(REF!|DIV/0!|VALUE!|NAME\?|N/A|NUM!|NULL!)", c.value):
                    errs += 1
    print(f"Сверка: расхождений {len(bad)}, ошибок формул {errs}")
    return not bad and not errs


# ============================================================================ main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--template", required=True, help="прошлая версия отчёта (.xlsx)")
    ap.add_argument("--data", required=True, help="выгрузка all_funds_perform_*.xlsx")
    ap.add_argument("--out", help="куда сохранить (по умолчанию funds_vs_benchmarks_<дата из имени выгрузки>.xlsx)")
    ap.add_argument("--cny", help="CSV курса ЦБ CNY/RUB: effective_date;rate (если cbr.ru недоступен)")
    ap.add_argument("--recalc", action="store_true", help="пересчитать формулы через LibreOffice и сверить значения")
    args = ap.parse_args()
    if not args.out:
        m = re.search(r"(\d{8})", Path(args.data).stem)
        args.out = str(Path(args.data).with_name(f"funds_vs_benchmarks_{m.group(1) if m else 'new'}.xlsx"))
    if Path(args.out).resolve() == Path(args.template).resolve():
        raise SystemExit("--out совпадает с --template")

    tmpl_data = pd.read_excel(args.template, sheet_name="Данные").set_index("Дата")
    lv, info = build_levels(args.data, tmpl_data, args.cny)
    meta, ctx = build_workbook(args, lv, info, tmpl_data)
    transplant_charts(args.template, args.out, args.out, meta)
    sc = ctx["scores"]
    print(f"Готово: {args.out}\nДанные {dt(lv.index.min())} – {dt(lv.index.max())}, {len(lv)} строк; "
          f"новых курсов ЦБ: {info['cny_new']}")
    for i, s in enumerate(sc):
        if s:
            mark = f"  → лист «{ctx['top_names'][i]}»" if i in ctx["top3"] else ""
            print(f"  {s['label']:10} обогнали {s['k']}/{s['n']} ({s['share']:.0%}), средняя альфа {dot_pct(s['mean'])}, "
                  f"медиана {dot_pct(s['median'])}, ранг {s['composite']:.2f}{mark}")
    if args.recalc and recalc(args.out):
        if not check(args.out, ctx):
            sys.exit(1)


if __name__ == "__main__":
    main()
