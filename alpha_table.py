#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Таблица «Альфа по фондам и периодам» из произвольного DataFrame — отдельным xlsx.

Вход — DataFrame, колонки которого идут парами «фонд, его бенчмарк»:

    moment | Акции | MCFTR | Облигации. РФ | RUCBTRNS | Ликвидность | RUSFAR | …

Названия колонок и задают таблицу: первая колонка пары — название фонда (столбец «Фонд»),
вторая — название бенчмарка под ним (столбец «Бенчмарк»). Ничего больше не нужно: списка
фондов в коде нет. Дубли названий допустимы (один бенчмарк у нескольких фондов) — пары
разбираются по позициям колонок, а не по именам. Число колонок с данными должно быть чётным.

Даты — в колонке «moment» (или «Дата»/«date»), в DatetimeIndex либо в первой колонке.
Значения — либо накопленные доходности в долях (0 на базовую дату, −0.012 = −1,2%), либо
уровни (цена пая, значение индекса, 1 + доходность). Тип определяется автоматически и
печатается в консоль. Один случай автоматика различить не может: срез доходностей, где все
значения положительны и первое далеко от нуля (например +85%) — выглядит как уровни. Для
такого среза передайте levels=False (или перебазируйте срез на его первую дату).

Из своего кода:

    from alpha_table import make_alpha_table

    df = my_func(...)                    # moment + пары «фонд, бенчмарк»
    path = make_alpha_table(df, "alpha_by_fund_30092026.xlsx")

Из терминала:

    python ./alpha_table.py 'path_to_excel' [out.xlsx] [лист]

Результат: один лист «Сводка» — таблица (шапка в A1, строка на пару, периоды 1 мес … 5 лет,
«Периодов с альфой > 0») и строка-пояснение двумя строками ниже. Валюта не учитывается.

Методика: прирост за период = уровень на конец / уровень на начало − 1, альфа = прирост фонда −
прирост бенчмарка. Обе границы периода — одна и та же дата для фонда и для бенчмарка: ближайшая
к календарной границе дата, где есть оба значения (в любую сторону, при равном удалении — более
ранняя), но не дальше 7 дней (near_days) и не дальше половины длины самого периода. Если такой
даты нет — в клетке «—»: у этого фонда нет истории на такой период.
Периоды: 1 и 2 недели, 1, 3, 6 месяцев, «с начала года», 1 год, 1,5 года (365 + 182 = 547 дней),
2, 3 и 5 лет. База «с начала года» — 30 декабря предыдущего года относительно последней даты
данных (выгрузка за 2026-й считается от 30.12.2025, за 2027-й — от 30.12.2026).
Модуль самодостаточный: нужны только pandas, numpy и openpyxl.
"""

from __future__ import annotations

import re
import sys
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pandas.tseries.offsets import DateOffset

__all__ = [
    "make_alpha_table",
    "alpha_dataframe",
    "prepare_levels",
    "bench_label",
    "ytd_start",
    "PERIODS",
]

# ============================================================================ настройки
YTD_LABEL = "С начала года"


def ytd_start(end: pd.Timestamp) -> pd.Timestamp:
    """База для «с начала года»: 30 декабря предыдущего года — последний торговый день года на MOEX.

    Год берётся от последней даты данных, а не от системного «сегодня»: пересчёт старой выгрузки
    тогда даёт правильную базу (данные до марта 2027 → 30.12.2026), а не дату из будущего.
    Если самого 30.12 в данных нет (суббота, как в 2028-м), границу подхватит обычный поиск
    ближайшей общей даты — возьмёт 29.12.
    """
    return pd.Timestamp(year=end.year - 1, month=12, day=30)


# Периоды: (подпись, сдвиг назад от последней даты ЛИБО функция «последняя дата → начало периода»).
# Полтора года — ровно 365 + 182 дня.
PERIODS = [
    ("1 неделя", pd.Timedelta(days=7)),
    ("2 недели", pd.Timedelta(days=14)),
    ("1 месяц", DateOffset(months=1)),
    ("3 месяца", DateOffset(months=3)),
    ("6 месяцев", DateOffset(months=6)),
    (YTD_LABEL, ytd_start),
    ("1 год", DateOffset(months=12)),
    ("1,5 года", pd.Timedelta(days=365 + 182)),
    ("2 года", DateOffset(months=24)),
    ("3 года", DateOffset(months=36)),
    ("5 лет", DateOffset(months=60)),
]
NEAR_DAYS = 7  # допуск у границы периода: ближайшая дата в любую сторону, ±7 дней
DATE_COLS = ("moment", "Дата", "date", "Date")

BASE_COLS = ["Фонд", "Бенчмарк"]
COUNT_COL = "Периодов с альфой > 0"
NOTE_COL = "Замечание"
DASH = "—"

# оформление как на листе «Сводка»
ARIAL = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="FF3A3A38")
ZEBRA_FILL = PatternFill("solid", fgColor="FFF2F2EF")
THIN = Side(style="thin", color="FFD6D5CE")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
GREY = "FF52514E"
# ширины в единицах Excel: px = width * 7 + 5 (65 → 460 px, 35 → 250 px)
FUND_WIDTH, BENCH_WIDTH, PERIOD_WIDTH, COUNT_WIDTH, NOTE_WIDTH = 65, 35, 11, 14, 75


# ============================================================================ данные
def _as_dates(values) -> pd.DatetimeIndex | None:
    """Даты, если значения на них похожи (≥ 90% распознаётся), иначе None."""
    s = pd.Series(values)
    if pd.api.types.is_numeric_dtype(s) and not pd.api.types.is_datetime64_any_dtype(s):
        return None  # RangeIndex / числа — не даты
    d = pd.to_datetime(s, dayfirst=True, errors="coerce")
    return pd.DatetimeIndex(d) if len(d) and d.notna().mean() >= 0.9 else None


def prepare_levels(
    data: pd.DataFrame,
    levels: bool | None = None,
    splits: list[tuple[str, str, float]] = (),
) -> pd.DataFrame:
    """DataFrame → уровни рядов (1 + накопленная доходность либо исходные цены) с датами в индексе.

    levels=None — определить автоматически (≈0 на старте → доходности, иначе уровни);
    splits — дробления паёв: (название колонки, дата, множитель), правятся только если
    скачок действительно есть в данных.
    """
    df = data.copy()
    idx = df.index if isinstance(df.index, pd.DatetimeIndex) else None
    if idx is None:
        col = next((c for c in DATE_COLS if c in df.columns), None)
        if col is not None:
            idx, df = _as_dates(df[col]), df.drop(columns=col)
        elif (idx := _as_dates(df.index)) is None:
            idx, df = _as_dates(df.iloc[:, 0]), df.iloc[:, 1:]  # даты в первой колонке
        if idx is None:
            raise ValueError(
                "Не нашёл даты: нужна колонка moment (или Дата/date), DatetimeIndex "
                "либо даты в первой колонке"
            )
    df.index = pd.DatetimeIndex(idx)
    df = df[df.index.notna()]
    df = df[~df.index.duplicated(keep="last")].sort_index()
    df = df.apply(pd.to_numeric, errors="coerce")
    first = [df.iloc[:, i].dropna() for i in range(df.shape[1])]
    start = (
        np.nanmedian([s.iloc[0] for s in first if len(s)])
        if any(len(s) for s in first)
        else np.nan
    )
    if np.isnan(start):
        raise ValueError("В колонках нет числовых данных")
    # уровни (цена пая, значение индекса) всегда > 0; доходности бывают отрицательными и стартуют с ≈0
    looks_returns = bool((df < 0).any().any()) or -0.5 < start < 0.5
    if levels is None:
        levels = not looks_returns
        print(
            f"Данные распознаны как {'уровни' if levels else 'накопленные доходности'} "
            f"(первое значение ≈ {start:.4g}); переопределяется флагом levels="
        )
    elif (
        levels is looks_returns
    ):  # явный флаг противоречит данным — предупредим, но послушаемся
        print(
            f"ВНИМАНИЕ: ряды похожи на {'доходности' if looks_returns else 'уровни'} "
            f"(первое значение ≈ {start:.4g}), а передано levels={levels}"
        )
    lv = df if levels else 1.0 + df
    for name, d, k in splits:
        d = pd.Timestamp(d)
        if not (lv.index.min() < d <= lv.index.max()):
            continue
        rows = np.flatnonzero(lv.index >= d)
        for pos in [i for i, c in enumerate(lv.columns) if c == name]:
            s = lv.iloc[:, pos]
            before, after = (
                s.loc[: d - pd.Timedelta(days=1)].dropna(),
                s.loc[d:].dropna(),
            )
            if len(before) and len(after) and after.iloc[0] / before.iloc[-1] < 2.0 / k:
                lv.iloc[rows, pos] = s.iloc[rows] * k
    x = lv.ffill()
    jumps = (x / x.shift(1) - 1).abs()
    big = [
        (lv.columns[i], d.date())
        for i in range(lv.shape[1])
        for d in jumps.index[jumps.iloc[:, i] > 0.5]
    ]
    if big:
        tail = f" … всего {len(big)}" if len(big) > 10 else ""
        print(
            "ВНИМАНИЕ: дневные изменения > 50% (дробление паёв?):",
            big[:10],
            tail,
            sep="",
        )
    return lv


def bench_label(bench: str, fund: str) -> str:
    """Название бенчмарка без приклеенного к нему названия фонда.

    В выгрузке колонка бенчмарка называется «<бенчмарк>, <название колонки фонда>», например
    'MCFTR, 8 | ОПИФРФИ "ВИМ - Акции"' → 'MCFTR', а 'БЕНЧ: БПИФРФИ "Ликвидность", 26 | …' →
    'БЕНЧ: БПИФРФИ "Ликвидность"'. Хвост отрезаем по точному совпадению с названием колонки
    фонда, и только если его нет — по первой запятой. Нет ни того, ни другого — отдаём как есть.
    """
    b, f = str(bench).strip(), str(fund).strip()
    if len(b) > len(f) and b.endswith(f):
        b = b[: -len(f)].strip().rstrip(",").strip()
    elif "," in b:
        b = b.split(",", 1)[0].strip()
    return b or str(bench).strip()


def fund_pairs(lv: pd.DataFrame) -> list[tuple[str, str, pd.Series, pd.Series]]:
    """Колонки по порядку → пары (название фонда, название бенчмарка, ряд фонда, ряд бенчмарка)."""
    names = [str(c) for c in lv.columns]
    if len(names) < 2 or len(names) % 2:
        raise ValueError(
            f"Нужно чётное число колонок «фонд + бенчмарк», получено {len(names)}: {names}"
        )
    return [
        (names[i], names[i + 1], lv.iloc[:, i], lv.iloc[:, i + 1])
        for i in range(0, len(names), 2)
    ]


# ============================================================================ расчёт
@dataclass
class PeriodDef:
    label: str
    nominal_start: pd.Timestamp  # граница периода «по календарю»
    end: pd.Timestamp  # последняя дата данных
    days: int  # номинальная длина периода в днях


def period_defs(dates: pd.DatetimeIndex) -> list[PeriodDef]:
    end = dates.max()
    out = []
    for label, spec in PERIODS:
        nom = spec(end) if callable(spec) else end - spec     # функция (YTD) или готовый сдвиг
        out.append(PeriodDef(label, nom, end, (end - nom).days))
    return out


def _near_pos(
    ok: np.ndarray, anchor: pd.Timestamp, dates: pd.DatetimeIndex, max_days: int
) -> int | None:
    """Позиция ближайшей к anchor даты, помеченной в ok; не дальше max_days дней, иначе None.

    ok — маска «на этой дате есть и фонд, и бенчмарк». Ищем в обе стороны: назад — когда на саму
    границу стоимости нет (пропуск, праздник), вперёд — когда история началась чуть позже границы.
    При равном удалении берётся более ранняя дата. anchor может лежать вне индекса.
    """
    n = len(dates)
    back = (
        int(dates.searchsorted(anchor, side="right")) - 1
    )  # последняя дата ≤ anchor, может быть -1
    fwd = back + 1
    while back >= 0 or fwd < n:
        db = (anchor - dates[back]).days if back >= 0 else None
        dw = (dates[fwd] - anchor).days if fwd < n else None
        if db is not None and db > max_days:
            back, db = -1, None  # назад вышли за окно
        if dw is not None and dw > max_days:
            fwd, dw = n, None
        if db is None and dw is None:
            return None
        if dw is None or (
            db is not None and db <= dw
        ):  # ближе назад (при равенстве — тоже назад)
            if ok[back]:
                return back
            back -= 1
        else:
            if ok[fwd]:
                return fwd
            fwd += 1
    return None


def period_alpha(
    f: pd.Series, b: pd.Series, p: PeriodDef, near_days: int = NEAR_DAYS
) -> float | None:
    """Альфа (ариф.) фонда к бенчмарку за период; None — у этой пары нет данных на период.

    Обе границы — ОДНА И ТА ЖЕ дата для фонда и для бенчмарка (ближайшая к границе, где есть оба).
    Иначе на пропуске в одном из рядов окна разъезжаются и альфа ловит движение рынка за эти дни.
    Допуск у границы — near_days, но не больше половины длины самого периода: для «1 недели» это
    3 дня, иначе семидневная натяжка превратила бы недельную альфу в двухдневную или двухнедельную.
    """
    ok = (f.notna() & b.notna()).values
    if not ok.any():
        return None
    dates, allow = f.index, min(near_days, max(1, p.days // 2))
    sp = _near_pos(ok, p.nominal_start, dates, allow)
    ep = _near_pos(ok, p.end, dates, allow)
    if sp is None or ep is None or ep <= sp:
        return None
    return (f.iat[ep] / f.iat[sp] - 1) - (b.iat[ep] / b.iat[sp] - 1)


def alpha_dataframe(
    data: pd.DataFrame,
    *,
    levels: bool | None = None,
    splits: list[tuple[str, str, float]] = (),
    notes: dict[str, str] | None = None,
    near_days: int = NEAR_DAYS,
) -> tuple[pd.DataFrame, list[PeriodDef]]:
    """Таблица «Альфа по фондам и периодам» как DataFrame (строка на пару «фонд + бенчмарк»)."""
    lv = prepare_levels(data, levels=levels, splits=splits)
    pairs = fund_pairs(lv)
    pdefs = period_defs(lv.index)
    cols = (
        BASE_COLS
        + [lbl for lbl, _ in PERIODS]
        + [COUNT_COL]
        + ([NOTE_COL] if notes else [])
    )
    rows = []
    for fname, bname, f, b in pairs:
        alphas = [period_alpha(f, b, p, near_days) for p in pdefs]
        nums = [a for a in alphas if a is not None]
        row = [
            fname,
            bench_label(bname, fname),
            *[DASH if a is None else a for a in alphas],
            f"{sum(a > 0 for a in nums)} из {len(nums)}",
        ]
        if notes:
            row.append(notes.get(fname))
        rows.append(row)
    return pd.DataFrame(rows, columns=cols), pdefs


# ============================================================================ xlsx
def _save_table(
    df: pd.DataFrame,
    pdefs: list[PeriodDef],
    out: Path,
    source: str | None,
    near_days: int = NEAR_DAYS,
) -> Path:
    cols = list(df.columns)
    wb = Workbook()
    ws = wb.active
    ws.title = "Сводка"
    ws.sheet_view.showGridLines = False
    first_p = len(BASE_COLS) + 1  # C
    last_p = first_p + len(PERIODS) - 1
    count_col = last_p + 1
    note_col = len(cols) if NOTE_COL in cols else None
    last_row = len(df) + 1

    for c, h in enumerate(cols, 1):
        cell = ws.cell(1, c, h)
        cell.font = Font(name=ARIAL, sz=9, b=True, color="FFFFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    ws.row_dimensions[1].height = 31.5

    for i, row in enumerate(df.itertuples(index=False), 2):
        zebra = i % 2 == 1  # как в оригинале: каждая вторая строка
        for c, v in enumerate(row, 1):
            cell = ws.cell(i, c)
            if c == count_col:  # «Периодов с альфой > 0» — формула от строки
                a, b = get_column_letter(first_p), get_column_letter(last_p)
                cell.value = f'=COUNTIF({a}{i}:{b}{i},">0")&" из "&COUNT({a}{i}:{b}{i})'
                cell.font = Font(name=ARIAL, sz=10)
            elif first_p <= c <= last_p and v != DASH:
                cell.value = float(v)
                cell.number_format = "0.00%"
                cell.font = Font(name=ARIAL, sz=10)
            elif v == DASH or c == note_col:
                cell.value = v if isinstance(v, str) else None
                cell.font = Font(name=ARIAL, sz=9, color=GREY)
                if v == DASH:
                    cell.alignment = Alignment(horizontal="center")
            else:
                cell.value = v
                cell.font = Font(name=ARIAL, sz=10)
                if c == 2:  # «Бенчмарк» — по центру
                    cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BORDER
            if zebra:
                cell.fill = ZEBRA_FILL

    rng = f"{get_column_letter(first_p)}2:{get_column_letter(last_p)}{last_row}"
    tl = f"{get_column_letter(first_p)}2"
    ws.conditional_formatting.add(
        rng,
        FormulaRule(
            formula=[f"AND(ISNUMBER({tl}),{tl}>0)"],
            font=Font(name=ARIAL, sz=10, b=True, color="FF1F5FA9"),
        ),
    )
    ws.conditional_formatting.add(
        rng,
        FormulaRule(
            formula=[f"AND(ISNUMBER({tl}),{tl}<0)"],
            font=Font(name=ARIAL, sz=10, color="FFB3372F"),
        ),
    )

    ws.column_dimensions["A"].width = FUND_WIDTH
    ws.column_dimensions["B"].width = BENCH_WIDTH
    for c in range(first_p, last_p + 1):
        ws.column_dimensions[get_column_letter(c)].width = PERIOD_WIDTH
    ws.column_dimensions[get_column_letter(count_col)].width = COUNT_WIDTH
    if note_col:
        ws.column_dimensions[get_column_letter(note_col)].width = NOTE_WIDTH
    ws.freeze_panes = "B2"
    ws.page_setup.orientation = "landscape"  # печать: одна страница в ширину
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    end = pdefs[0].end
    ytd = next((x for x in pdefs if x.label == YTD_LABEL), None)
    note = (
        f"Альфа = прирост фонда − прирост бенчмарка за период (арифметическая). Конец всех периодов — "
        f"{end:%d.%m.%Y}. «1,5 года» — 547 дней (365 + 182). "
    )
    if ytd is not None:
        note += f"«{YTD_LABEL}» — от {ytd.nominal_start:%d.%m.%Y}. "
    note += (
        f"Границы периода — одна и та же дата для фонда и бенчмарка: ближайшая к границе, где есть "
        f"оба значения, но не дальше {near_days} дней (для коротких периодов — не дальше половины "
        f"их длины). Синий — фонд обогнал бенчмарк, красный — отстал, «—» — у фонда нет данных "
        f"на этот период."
    )
    if source:
        note += f" Источник: {source}."
    c = ws.cell(last_row + 2, 1, note)
    c.font = Font(name=ARIAL, sz=9, color=GREY)

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def make_alpha_table(
    data: pd.DataFrame,
    out: str | Path | None = None,
    *,
    levels: bool | None = None,
    splits: list[tuple[str, str, float]] = (),
    notes: dict[str, str] | None = None,
    near_days: int = NEAR_DAYS,
    source: str | None = None,
) -> Path:
    """DataFrame → xlsx с таблицей «Альфа по фондам и периодам». Возвращает путь к файлу.

    data   — даты (колонка moment / индекс / первая колонка) + колонки парами «фонд, бенчмарк»;
    out    — куда сохранить; по умолчанию alpha_by_fund_<ДДММГГГГ последней даты>.xlsx в текущей папке;
    levels — None: определить автоматически; True — в data уровни (цены/индексы/1+доходность),
             False — накопленные доходности в долях;
    splits — дробления паёв: [(название колонки, 'ГГГГ-ММ-ДД', множитель)];
    notes  — замечания по фондам {название колонки фонда: текст}; столбец «Замечание» появляется,
             только если словарь передан;
    near_days — допуск поиска значения у границы периода, дней (по умолчанию 7): нет стоимости на
             нужную дату — берётся ближайшая в любую сторону внутри этого окна;
    source — подпись источника в строке-пояснении под таблицей.
    """
    if not isinstance(data, pd.DataFrame):
        raise TypeError(f"Ожидался pandas.DataFrame, получен {type(data).__name__}")
    table, pdefs = alpha_dataframe(
        data, levels=levels, splits=splits, notes=notes, near_days=near_days
    )
    if out is None:
        out = Path.cwd() / f"alpha_by_fund_{pdefs[0].end:%d%m%Y}.xlsx"
    return _save_table(table, pdefs, Path(out), source, near_days)


# ============================================================================ запуск из терминала
def main(argv: list[str]) -> Path:
    if not argv:
        raise SystemExit(
            "Запуск: python ./alpha_table.py 'path_to_excel' [out.xlsx] [лист]"
        )
    src = Path(argv[0]).expanduser().resolve()
    if not src.exists():
        raise SystemExit(f"Нет файла {src}")
    xl = pd.ExcelFile(src)
    sheet = argv[2] if len(argv) > 2 else xl.sheet_names[0]
    data = xl.parse(sheet)
    data = data[
        [
            c
            for c in data.columns
            if not (isinstance(c, str) and c.startswith("Unnamed"))
        ]
    ]
    if len(argv) > 1:
        out = Path(argv[1])
    else:
        m = re.search(r"(\d{8})", src.stem)
        out = src.with_name(f"alpha_by_fund_{m.group(1)}.xlsx") if m else None
    res = make_alpha_table(data, out, source=f"лист «{sheet}» файла {src.name}")
    print(f"Готово: {res}")
    return res


if __name__ == "__main__":
    try:
        main(sys.argv[1:])
    except ValueError as e:
        raise SystemExit(f"Ошибка: {e}")
